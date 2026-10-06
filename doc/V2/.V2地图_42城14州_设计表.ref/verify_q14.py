# -*- coding: utf-8 -*-
"""校验 V2 地图设计表：两个 sheet 的行列与统计。"""

from collections import Counter, defaultdict

import openpyxl

PATH = r"I:/buddyWork/sanWalk/doc/V2/V2地图_42城14州_设计表.xlsx"

wb = openpyxl.load_workbook(PATH)
print("sheets:", wb.sheetnames)

ws1 = wb["地图设计表"]
print(f"[地图设计表] dims={ws1.dimensions} max_row={ws1.max_row} max_col={ws1.max_column}")
print("  header:", [c.value for c in ws1[1]])

ws2 = wb["设施清单"]
print(f"[设施清单] dims={ws2.dimensions} max_row={ws2.max_row} max_col={ws2.max_column}")
print("  header:", [c.value for c in ws2[1]])

rows = list(ws2.iter_rows(min_row=2, values_only=True))
print(f"  数据行数: {len(rows)}")

zhou = Counter(r[0] for r in rows)
seat = sum(1 for r in rows if r[2] == "是")
level = Counter(r[3] for r in rows)
port = sum(1 for r in rows if r[10] == "是")
primary = Counter(f"{r[4]}/{r[5]}" for r in rows)
secondary = Counter(f"{r[6]}/{r[7]}" for r in rows if r[6])

rare = defaultdict(list)
for r in rows:
    for m in str(r[9]).split("、"):
        if m:
            rare[m].append(r[1])

print("  州数:", len(zhou), dict(zhou))
print("  州府:", seat, "｜城级:", dict(level), "｜港口:", port)
print("  主设施分布:", dict(primary))
print("  副设施分布:", dict(secondary))
print("  稀有材料 → 城:")
for m, cities in rare.items():
    print(f"    {m}: {'、'.join(cities)} ({len(cities)} 城)")

assert len(rows) == 42 and len(wb.sheetnames) == 2
print("OK")
