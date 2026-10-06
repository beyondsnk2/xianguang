# -*- coding: utf-8 -*-
"""V2 地图设计表：42 城 · 14 州（州 / 城市 / 类型）"""

try:
    import openpyxl
except ImportError:
    import subprocess, sys
    subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", "openpyxl>=3.1.0"])
    import openpyxl

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation


def xl_color(css_hex: str) -> str:
    value = css_hex.removeprefix("#").upper()
    if len(value) != 6:
        raise ValueError(f"Expected #RRGGBB, got: {css_hex}")
    return "FF" + value


XL_HEADER_BG = xl_color("#4472C4")
XL_HEADER_FG = xl_color("#FFFFFF")
XL_SEAT_BG = xl_color("#D9E2F3")      # 州府行浅底
XL_BORDER = xl_color("#BFBFBF")

thin = Side(style="thin", color=XL_BORDER)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal="center", vertical="center")
HEADER_FONT = Font(bold=True, color=XL_HEADER_FG, size=11)
BODY_FONT = Font(size=11)
SEAT_FONT = Font(size=11, bold=True)
HEADER_FILL = PatternFill("solid", fgColor=XL_HEADER_BG)
SEAT_FILL = PatternFill("solid", fgColor=XL_SEAT_BG)

OUT_XLSX = r"I:/buddyWork/sanWalk/doc/V2/V2地图_42城14州_设计表.xlsx"
TITLE = "V2地图_42城14州_设计表"

# 列：州 / 城市 / 类型    （州序：幽并冀青司隶兖徐豫雍凉荆扬益交；州府置于各州首行）
ROWS = [
    ("幽州", "蓟",   "州府·中城"),
    ("幽州", "北平", "中城"),
    ("幽州", "襄平", "中城"),

    ("并州", "晋阳", "州府·中城"),
    ("并州", "上党", "小城"),

    ("冀州", "邺",   "州府·大城"),
    ("冀州", "南皮", "中城"),
    ("冀州", "平原", "小城"),

    ("青州", "北海", "州府·中城"),
    ("青州", "济南", "小城"),

    ("司隶", "洛阳", "都城"),
    ("司隶", "河内", "中城"),

    ("兖州", "濮阳", "州府·中城"),
    ("兖州", "陈留", "中城"),

    ("徐州", "下邳", "州府·中城"),
    ("徐州", "小沛", "小城"),
    ("徐州", "广陵", "中城"),

    ("豫州", "许昌", "州府·大城"),
    ("豫州", "汝南", "中城"),
    ("豫州", "谯",   "中城"),

    ("雍州", "长安", "州府·大城"),
    ("雍州", "天水", "中城"),
    ("雍州", "安定", "小城"),

    ("凉州", "武威", "州府·中城"),
    ("凉州", "西平", "小城"),

    ("荆州", "襄阳", "州府·大城"),
    ("荆州", "宛",   "中城"),
    ("荆州", "新野", "小城"),
    ("荆州", "江夏", "中城"),
    ("荆州", "江陵", "中城"),
    ("荆州", "长沙", "中城"),

    ("扬州", "建业", "州府·大城"),
    ("扬州", "吴",   "中城"),
    ("扬州", "会稽", "中城"),
    ("扬州", "柴桑", "中城"),
    ("扬州", "庐江", "小城"),

    ("益州", "成都", "州府·大城"),
    ("益州", "汉中", "中城"),
    ("益州", "梓潼", "小城"),
    ("益州", "江州", "中城"),

    ("交州", "交趾", "州府·中城"),
    ("交州", "南海", "中城"),
]

assert len(ROWS) == 42, f"行数应为 42，实际 {len(ROWS)}"

wb = Workbook()
ws = wb.active
ws.title = "地图设计表"
wb.properties.title = TITLE

# —— 锚点：第 1 行表头，第 2~43 行数据 ——
HEADER_ROW = 1
DATA_START = 2
DATA_END = DATA_START + len(ROWS) - 1  # 43

headers = ["州", "城市", "类型"]
for col, name in enumerate(headers, start=1):
    c = ws.cell(row=HEADER_ROW, column=col, value=name)
    c.font = HEADER_FONT
    c.fill = HEADER_FILL
    c.alignment = CENTER
    c.border = BORDER

for i, (zhou, city, ctype) in enumerate(ROWS):
    r = DATA_START + i
    is_seat = ctype.startswith("州府") or ctype == "都城"
    for col, val in enumerate((zhou, city, ctype), start=1):
        c = ws.cell(row=r, column=col, value=val)
        c.font = SEAT_FONT if (is_seat and col == 2) else BODY_FONT
        c.alignment = CENTER
        c.border = BORDER
        if is_seat:
            c.fill = SEAT_FILL

# 列宽
ws.column_dimensions["A"].width = 10
ws.column_dimensions["B"].width = 12
ws.column_dimensions["C"].width = 16

# 冻结表头 + 筛选（42 行明细）
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A{HEADER_ROW}:C{DATA_END}"

# 类型列数据验证（枚举）
dv = DataValidation(
    type="list",
    formula1='"都城,州府·大城,州府·中城,中城,小城"',
    allow_blank=True,
    showDropDown=False,
)
dv.error = "类型只能取：都城 / 州府·大城 / 州府·中城 / 中城 / 小城"
dv.errorTitle = "类型不合法"
ws.add_data_validation(dv)
dv.add(f"C{DATA_START}:C{DATA_END}")

wb.save(OUT_XLSX)
print(f"saved: {OUT_XLSX}")
print(f"rows: {len(ROWS)}  data range: A{HEADER_ROW}:C{DATA_END}")
