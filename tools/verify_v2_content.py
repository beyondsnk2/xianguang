"""校验 firstShow_V2.xlsx 的内容表是否满足六条设计规则。"""
from pathlib import Path
import openpyxl

SRC = Path(r"I:/buddyWork/sanWalk/config/firstShow_V2.xlsx")
wb = openpyxl.load_workbook(SRC, data_only=True)


def rows(title):
    ws = wb[title]
    hdr = [c.value for c in ws[1]]
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        out.append(dict(zip(hdr, r)))
    return hdr, out


ih, items = rows("item")
rh, recipes = rows("recipe")
bh, bps = rows("blueprint")
th, tasks = rows("taskTpl")
sh, skills = rows("skill")

errs = []
ok = []

# 索引
item_by_tag = {it["tag"]: it for it in items}

# ① 稀有材料只 C 类产
rare_src = {it["tag"]: it["source"] for it in items if it["cat"] == "稀有"}
bad = [t for t, s in rare_src.items() if s != "C类独占"]
if bad:
    errs.append(f"①稀有材料非C类产: {bad}")
else:
    ok.append(f"① 稀有材料({len(rare_src)}种)全部 source=C类独占")

# ② 图纸只 C 类给（blueprint.fromSkill 全为 C 类技能）
c_skills = {s["tag"] for s in skills if s["cls"] == "C"}
bad_bp = [b["tag"] for b in bps if b["fromSkill"] not in c_skills]
if bad_bp:
    errs.append(f"②图纸非C类技能产: {bad_bp}")
else:
    ok.append(f"② 图纸({len(bps)}张)来源技能全为C类: {[b['fromSkill'] for b in bps]}")

# ③ 配方三城互异（上品：craft/rare/blueprint 三城互不相同）
def cities_distinct(rec):
    cs = [rec["cityCraft"], rec["cityRare"], rec["cityBlueprint"]]
    cs = [c for c in cs if c]
    return len(cs) == len(set(cs)), cs

high_bad = []
for rec in recipes:
    if rec["quality"] >= 7 and rec["needBlueprint"]:
        d, cs = cities_distinct(rec)
        if not d:
            high_bad.append((rec["tag"], cs))
if high_bad:
    errs.append(f"③上品配方三城非互异: {high_bad}")
else:
    ok.append(f"③ 上品配方({sum(1 for r in recipes if r['quality']>=7)}) 三城(制/稀/图)均互异")

# ④ 投喂下界：n品制造只吃 q>=n 材料；且成品品质 >= matQualityFloor
feed_bad = []
for rec in recipes:
    mat = item_by_tag.get(rec["needItem1"])
    if not mat:
        feed_bad.append((rec["tag"], "needItem1缺失:" + str(rec["needItem1"])))
        continue
    # 材料品质下界
    mq = mat["qMin"]
    if mq < rec["matQualityFloor"]:
        feed_bad.append((rec["tag"], f"材料{mq}低于投喂下界{rec['matQualityFloor']}"))
    if rec["quality"] < rec["matQualityFloor"]:
        feed_bad.append((rec["tag"], f"成品品{rec['quality']}<投喂下界{rec['matQualityFloor']}"))
    # 稀有料门槛：品质>=4（中品起）才需稀有料；1-3 下品无料无图
    if rec["quality"] >= 4 and not rec["needRare"]:
        feed_bad.append((rec["tag"], f"品质{rec['quality']}≥4却无稀有料"))
    if rec["quality"] < 4 and rec["needRare"]:
        feed_bad.append((rec["tag"], f"品质{rec['quality']}<4却有稀有料"))
if feed_bad:
    errs.append("④投喂下界/稀有料门槛异常: " + str(feed_bad))
else:
    ok.append("④ 投喂下界 + 稀有料门槛(品质4起) 全部满足")

# ⑤ 稀有料产出品质下界 >= 4
badq = [t for t, it in item_by_tag.items() if it["cat"] == "稀有" and it["qMin"] < 4]
if badq:
    errs.append(f"⑤稀有料品质下界<4: {badq}")
else:
    ok.append(f"⑤ 稀有料品质下界均≥4 (qMin范围 {min(it['qMin'] for it in items if it['cat']=='稀有')}~{max(it['qMin'] for it in items if it['cat']=='稀有')})")

# ⑥ 同品咬合：同品B = 同品A料(量) + 同品C稀有料(质)
#   检查每条需要稀有料的配方：其稀有料必须由某C任务在对应城可产（城市名匹配 cityRare 在 RARE 城市列表中）
#   简化：核对 recipe.cityRare 与 item(rare).note 中的城市是否一致
rare_city = {}
for it in items:
    if it["cat"] == "稀有":
        # note 形如 "产出城:晋阳;邺"
        note = it["note"] or ""
        if "产出城:" in note:
            rare_city[it["tag"]] = set(note.split("产出城:")[1].split(";"))
rare_mismatch = []
for rec in recipes:
    # 同品咬合 tier 对齐：A 料档位 == 品质；稀有料档位 == 品质
    mat = item_by_tag.get(rec["needItem1"])
    if mat and mat["qMin"] != rec["quality"]:
        rare_mismatch.append((rec["tag"], f"A料档{mat['qMin']}≠品质{rec['quality']}"))
    if rec["needRare"]:
        rk = item_by_tag.get(rec["needRare"])
        if rk and rk["qMin"] != rec["quality"]:
            rare_mismatch.append((rec["tag"], f"稀有档{rk['qMin']}≠品质{rec['quality']}"))
        cities = rare_city.get(rec["needRare"], set())
        if rec["cityRare"] and rec["cityRare"] not in cities:
            rare_mismatch.append((rec["tag"], rec["needRare"], rec["cityRare"], sorted(cities)))
if rare_mismatch:
    errs.append("⑥同品咬合异常: " + str(rare_mismatch))
else:
    ok.append(f"⑥ 同品咬合：{sum(1 for r in recipes if r['needRare'])}条需稀有料配方，A料/稀有料档位均=品质，cityRare 均落在该料产出城内")

# 任务基型数（9 类 × 9 品质 = 81）
if len(tasks) != 81:
    errs.append(f"任务基型数={len(tasks)}，应为81")
else:
    ok.append(f"任务基型=81 (9类×9品质)")
if len(recipes) != 315:
    errs.append(f"配方数={len(recipes)}，应为315")
else:
    ok.append("配方=315 (35族×9品质)")

print("=" * 60)
for o in ok:
    print("✅", o)
if errs:
    print("-" * 60)
    for e in errs:
        print("❌", e)
    print("=" * 60)
    print("校验未通过")
else:
    print("=" * 60)
    print("全部六条校验通过 ✅")
