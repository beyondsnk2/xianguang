"""Plan A (2026-10-07) 直接改运行时 xlsx：B 类 q7-9 的 getItemNum 由 1;3 → 1;2。
只动 task 表的 B+quality>=7 行，不动 item 价格列 / 其它表。
"""
import openpyxl

SRC = "config/firstShow_V4.xlsx"

wb = openpyxl.load_workbook(SRC)
ws = wb["task"]

# header = row 1
hdr = [c.value for c in ws[1]]
tag_i = hdr.index("tag")
cls_i = hdr.index("cls")
q_i = hdr.index("quality")
num_i = hdr.index("getItemNum")

changed = []
for row in ws.iter_rows(min_row=2):
    tag = row[tag_i].value
    cls = row[cls_i].value
    q = row[q_i].value
    num = row[num_i].value
    if cls == "B" and isinstance(q, int) and q >= 7:
        if num == "1;3":
            row[num_i].value = "1;2"
            changed.append((tag, q, num, "1;2"))
        else:
            changed.append((tag, q, num, f"SKIP(已是 {num})"))

wb.save(SRC)
print("patched rows:")
for c in changed:
    print(" ", c)
print("total B q>=7 rows touched:", len(changed))
